import csv
import unittest
from pathlib import Path

import openpyxl


ROOT = Path(__file__).resolve().parents[1]
WORKBOOK = ROOT / "data" / "processed" / "10124-users.organized.xlsx"
CLEANED = ROOT / "data" / "processed" / "10124-users.cleaned.csv"
ISSUES = ROOT / "data" / "processed" / "10124-users.cleaning-issues.csv"


class OrganizedWorkbookTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.workbook = openpyxl.load_workbook(WORKBOOK, read_only=False, data_only=True)

    @classmethod
    def tearDownClass(cls):
        cls.workbook.close()

    def test_workbook_has_requested_sheets(self):
        self.assertEqual(
            self.workbook.sheetnames,
            [
                "Contacts",
                "Data Quality Notes",
                "Possible Duplicate People",
                "Audit Log",
            ],
        )

    def test_contacts_row_count_ids_and_employee_range(self):
        contacts = self.workbook["Contacts"]
        headers = [cell.value for cell in contacts[1]]
        self.assertEqual(contacts.max_row - 1, 47613)
        self.assertEqual(headers[headers.index("employees") + 1], "employee_range")

        with CLEANED.open(newline="", encoding="utf-8-sig") as handle:
            cleaned_first = next(csv.DictReader(handle))

        first_row = {
            header: contacts.cell(row=2, column=index + 1).value
            for index, header in enumerate(headers)
        }
        self.assertEqual(first_row["id"], cleaned_first["id"])
        self.assertEqual(first_row["employees"], cleaned_first["employees"])
        self.assertEqual(first_row["employee_range"], "10,001-50,000")

    def test_text_columns_are_text_formatted(self):
        contacts = self.workbook["Contacts"]
        headers = [cell.value for cell in contacts[1]]
        for column_name in [
            "desk_phone",
            "corporate_phone",
            "mobile_phone",
            "postal_code",
            "executive_postal_code",
        ]:
            column_index = headers.index(column_name) + 1
            self.assertEqual(contacts.cell(row=2, column=column_index).number_format, "@")

    def test_duplicate_sheet_counts(self):
        duplicates = self.workbook["Possible Duplicate People"]
        headers = [cell.value for cell in duplicates[1]]
        status_col = headers.index("duplicate_status") + 1
        group_col = headers.index("duplicate_group_id") + 1
        statuses = [
            duplicates.cell(row=row, column=status_col).value
            for row in range(2, duplicates.max_row + 1)
        ]
        groups = {
            duplicates.cell(row=row, column=group_col).value
            for row in range(2, duplicates.max_row + 1)
        }
        self.assertEqual(duplicates.max_row - 1, 1663)
        self.assertEqual(len(groups), 811)
        self.assertEqual(statuses.count("duplicate"), 852)
        self.assertEqual(statuses.count("primary"), 811)

    def test_audit_log_count_matches_issue_csv(self):
        audit_log = self.workbook["Audit Log"]
        with ISSUES.open(newline="", encoding="utf-8-sig") as handle:
            issue_count = sum(1 for _ in csv.DictReader(handle))
        self.assertEqual(audit_log.max_row - 1, issue_count)
        self.assertEqual(issue_count, 34483)


    def test_all_sheets_freeze_header_row(self):
        for sheet_name in self.workbook.sheetnames:
            self.assertEqual(self.workbook[sheet_name].freeze_panes, "A2")

if __name__ == "__main__":
    unittest.main()

